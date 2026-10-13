"""Authoring helper: faithful vector diagrams from the supplied PowerPoint.
Requires python-pptx only when regenerating assets; SCEN has no runtime dependency.
All focus and answer states use the SAME SVG coordinates as their diagram.
"""
from pathlib import Path
from pptx import Presentation
from lxml import etree as E
import html,json,base64,math,re,textwrap
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'bilder/minimax/v2';OUT.mkdir(exist_ok=True)
P=Presentation(ROOT/'docs/minimax/original.pptx');SCALE=144/914400
BG='#080f18';INK='#eef4f7';LINE='#8da1b5';MAX='#66ddbb';MIN='#f4ae66'
def esc(s):return html.escape(str(s))
def txt(x,y,s,size=30,c=INK,anchor='middle'):
 return f'<text x="{x:.2f}" y="{y:.2f}" font-family="Arial,sans-serif" font-size="{size}" fill="{c}" text-anchor="{anchor}">{esc(s)}</text>'
def save(name,body,box=(0,0,1920,1000)):
 x,y,w,h=box
 (OUT/(name+'.svg')).write_text(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{x} {y} {w} {h}" role="img"><rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{BG}"/>'+body+'</svg>')
def flatten(shapes,sx=SCALE,sy=SCALE,tx=0,ty=0,parent=None):
 out=[]
 for s in shapes:
  x=s.left*sx+tx;y=s.top*sy+ty;w=s.width*sx;h=s.height*sy
  el=s._element;st=s.text.strip() if s.has_text_frame else ''
  if s.shape_type==6:
   g={'id':s.shape_id,'type':'group','x':x,'y':y,'w':w,'h':h,'parent':parent};out.append(g)
   xf=el.find('{*}grpSpPr/{*}xfrm');off=xf.find('{*}chOff');ex=xf.find('{*}chExt')
   out.extend(flatten(s.shapes,w/int(ex.get('cx')),h/int(ex.get('cy')),x-int(off.get('x'))*w/int(ex.get('cx')),y-int(off.get('y'))*h/int(ex.get('cy')),s.shape_id));continue
  xf=el.find('{*}spPr/{*}xfrm');geom=el.find('{*}spPr/{*}prstGeom');g=geom.get('prst') if geom is not None else ''
  d={'id':s.shape_id,'type':'picture' if s.shape_type==13 else 'line' if s.shape_type==9 else 'shape','geom':g,'x':x,'y':y,'w':w,'h':h,'text':st,'rot':float(xf.get('rot','0'))/60000 if xf is not None else 0,'fh':xf is not None and xf.get('flipH')=='1','fv':xf is not None and xf.get('flipV')=='1','parent':parent}
  if d['type']=='picture':d['image']=f'data:{s.image.content_type};base64,{base64.b64encode(s.image.blob).decode()}'
  out.append(d)
 return out
DATA={n:flatten(P.slides[n-1].shapes) for n in [6,7,8,10,13,20]}
def source_render(n,visible=None,active=(),shade_ids=(),omit=(),box=(0,0,1920,1000)):
 ds=DATA[n];byid={d['id']:d for d in ds};body='';seen=[]
 def focus(d):
  x=d['x'];y=d['y'];w=d['w'];h=d['h'];return f'<rect x="{x-14:.2f}" y="{y-14:.2f}" width="{w+28:.2f}" height="{h+28:.2f}" rx="10" fill="{MAX}" fill-opacity=".10" stroke="{MAX}" stroke-width="5" data-focus-id="{d["id"]}"/>'
 for key in active:
  if key in byid:body+=focus(byid[key])
 for d in ds:
  if d['id'] in omit:continue
  x,y,w,h=[d[k] for k in ['x','y','w','h']];st=d.get('text','');g=d.get('geom','');par=byid.get(d.get('parent'))
  if d['type']=='group':continue
  if 'Presentation av David' in st or st.startswith(('Beslutsträd för','Uppgift','Aktivitet')):continue
  if visible is not None and st and re.match(r'^(?:Min-värde|Max-värde|Värde|[−-]?\d+$)',st) and d['id'] not in visible:continue
  alpha='.65' if d['id'] in shade_ids else '1';s=''
  if d['type']=='picture':s=f'<image x="{x}" y="{y}" width="{w}" height="{h}" href="{d["image"]}"/>'
  elif d['type']=='line':
   color='#f6c27c' if par and par['w']<80 and par['h']<80 else LINE
   s=f'<line x1="{x+w if d["fh"] else x}" y1="{y+h if d["fv"] else y}" x2="{x if d["fh"] else x+w}" y2="{y if d["fv"] else y+h}" stroke="{color}" stroke-width="3" stroke-linecap="round"/>'
  elif g=='ellipse':s=f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" fill="none" stroke="{MAX}" stroke-width="3.5"/>'
  elif g=='triangle':
   inv=abs(d['rot']%360-180)<1;color=MIN if inv else MAX;pts=f'{x},{y} {x+w},{y} {x+w/2},{y+h}' if inv else f'{x+w/2},{y} {x+w},{y+h} {x},{y+h}'
   s=f'<polygon points="{pts}" fill="{color}" fill-opacity=".14" stroke="{color}" stroke-width="3"/>'
  elif g in ['rect','roundRect'] and not st:
   # Small rectangles next to boards are original answer labels; retain them.
   s=f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{10 if g=="roundRect" else 0}" fill="none" stroke="{LINE}" stroke-width="2.5"/>'
  if st:
   # Numeric labels align with their containing triangle, not the textbox's margin.
   duplicates=any(old[0]==st and abs(old[1]-x)<8 and abs(old[2]-y)<8 for old in seen)
   if duplicates:continue
   seen.append((st,x,y))
   fs=28 if n in [8,10,13,20] else 31
   if st.startswith('”Jag baserar'):fs=29
   if re.fullmatch(r'-?\d+',st):
    triangles=[t for t in ds if t.get('geom')=='triangle'];nearest=min(triangles,key=lambda t:(t['x']+t['w']/2-x-w/2)**2+(t['y']+t['h']/2-y-h/2)**2) if triangles else None
    if nearest and abs(nearest['x']+nearest['w']/2-x-w/2)<100 and abs(nearest['y']+nearest['h']/2-y-h/2)<80:
     fs=31;s+=txt(nearest['x']+nearest['w']/2,nearest['y']+nearest['h']/2+fs*.28,st,fs)
    else:s+=txt(x+w/2,y+h/2+fs*.28,st,fs)
   else:
    # Preserve original labels as one line where possible.
    width=max(w,260) if st.startswith(('Max-värde','Min-värde','Värde')) else w
    for j,line in enumerate(textwrap.wrap(st,max(12,int(width/(fs*.49))))):s+=txt(x+6,y+fs+j*fs*1.22,line,fs,anchor='start')
  if s:body+=f'<g data-source-id="{d["id"]}" opacity="{alpha}">{s}</g>'
 return body
# Full source diagram geometry. Frames differ only in their focus or revealed labels.
for i,active in enumerate([(),(3,),(9,15,21,27,33,39,45,51,57),(3,9)],0):
 # Find board groups from width, rather than hardcoded list ids from a different deck.
 ds=DATA[6];boards=[d for d in ds if d['type']=='group' and d['w']>120 and d['h']>120]
 active=() if i==0 else (boards[0]['id'],) if i==1 else tuple(d['id'] for d in boards[1:10]) if i==2 else (boards[1]['id'],boards[-1]['id'])
 save(f'speltrad-{i}',source_render(6,active=active),box=(0,75,1920,985))
for n,labelsets in [(7,[[132,133],[130,131],[129]]),(8,[[296,297],[302,303,304],[300,301,305,306],[298],[307]]),(10,[[397,414,415,417],[419,420,422,423,424,426],[427,428,429],[430]])]:
 # Shape ids in slide7 are one greater than names printed by Office.
 ds=DATA[n];alllabels=[d for d in ds if re.match(r'^(?:Min-värde|Max-värde|Värde|[−-]?\d+$)',d.get('text',''))]
 if n==7:
  labelsets=[[d['id'] for d in alllabels if d['text'].startswith('Värde')],[d['id'] for d in alllabels if d['text'].startswith('Max-värde')],[d['id'] for d in alllabels if d['text'].startswith('Min-värde')]]
 visible=set();save(f'spel-{n}-0',source_render(n,visible),box=(0,0,1920,1015))
 for i,ids in enumerate(labelsets,1):
  visible.update(ids);save(f'spel-{n}-{i}',source_render(n,visible,active=ids),box=(0,0,1920,1015))
# Large original exercise: restore all topology and reveal answers level by level.
n=13;labels=[d for d in DATA[n] if re.fullmatch(r'\d+',d.get('text',''))];leaf=[d['id'] for d in labels if d['y']>910];vis=set(leaf);save('rotnod-0',source_render(n,vis),box=(0,210,1920,835))
for i,(lo,hi) in enumerate([(740,910),(530,740),(370,530),(0,370)],1):
 ids=[d['id'] for d in labels if lo<=d['y']<hi];vis.update(ids);save(f'rotnod-{i}',source_render(n,vis,active=ids),box=(0,210,1920,835))
# Exact numerical models preserve the original up/down triangle convention.
def tree(name,nodes,edges,values,active=(),pruned=(),bounds=None,count='',alpha=False):
 body=txt(100,34,'▲ MAX',25,MAX,'start')+txt(260,34,'▼ MIN',25,MIN,'start')
 size=.75 if len(nodes)>18 else 1
 def vertices(k):
  x,y,role=nodes[k];dx=62*size;dy=55*size
  return [(x-dx,y-dy),(x+dx,y-dy),(x,y+dy)] if role=='min' else [(x,y-dy),(x+dx,y+dy),(x-dx,y+dy)]
 def boundary(k,other):
  x,y,_=nodes[k];_,yy,_=nodes[other]
  # Trees have one input and one output port; lines stop at the triangle edge.
  return x,y+(55*size if yy>y else -55*size)

 for a,b in edges:
  x,y,role=nodes[a];xx,yy,_=nodes[b];cut=b in pruned;color=MAX if b in active else LINE
  x,y=boundary(a,b);xx,yy=boundary(b,a)
  body+=f'<path data-edge="{a}:{b}" d="M{x},{y}L{xx},{yy}" stroke="{color}" stroke-width="{5 if b in active else 2.5}" fill="none" opacity="{.45 if cut else .85}"'+(' stroke-dasharray="8 7"' if cut else '')+'/>'
  if cut:
   mx=x+(xx-x)*.64;my=y+(yy-y)*.64
   body+=f'<path d="M{mx-14},{my-14}L{mx+14},{my+14}M{mx+14},{my-14}L{mx-14},{my+14}" stroke="#ff8581" stroke-width="5"/>'
 for key,(x,y,role) in nodes.items():
  color=MIN if role=='min' else MAX;pts=' '.join(f'{xx},{yy}' for xx,yy in vertices(key))
  if key in active:body+=f'<rect data-focus-node="{key}" x="{x-82}" y="{y-72}" width="164" height="144" rx="12" fill="{color}" fill-opacity=".09" stroke="{color}" stroke-width="4"/>'
  body+=f'<g data-node="{key}" opacity="{.45 if key in pruned else 1}"><polygon points="{pts}" fill="{color}" fill-opacity=".16" stroke="{color}" stroke-width="3"/>'+txt(x,y+12,values.get(key,''),32 if size<1 else 40)+ '</g>'
  if bounds and key in bounds:
   a,b=bounds[key];body+=txt(x+92,y+10,f'α = {a}    β = {b}',26,INK,'start')
 save(name,body,box=(0,0,1800,810))
node={'r':(900,130,'max'),'a':(360,360,'min'),'b':(900,360,'min'),'c':(1440,360,'min')};edges=[('r',k) for k in 'abc'];lk=[]
for parent,x in zip('abc',[360,900,1440]):
 for j,dx in enumerate([-155,0,155]):k=parent+str(j);node[k]=(x+dx,660,'max');edges.append((parent,k));lk.append(k)
leaves=dict(zip(lk,[8,5,6,9,3,7,2,4,6]));v=leaves.copy();tree('minimax-0',node,edges,v)
for i,(k,answer) in enumerate(zip('abc',[5,3,2]),1):v[k]=answer;tree(f'minimax-{i}',node,edges,v,[k]+[k+str(j) for j in range(3)])
v['r']=5;tree('minimax-4',node,edges,v,['r','a'])
# Conceptual pruning: original known leaf scores remain visible, while irrelevant branches fade.
v=leaves.copy();v['a']=5;v['r']=5;tree('beskarning-0',node,edges,v,['a','r'])
v['b']='≤3';tree('beskarning-1',node,edges,v,['b','b1'],['b2'])
v['c']='≤2';tree('beskarning-2',node,edges,v,['c','c0'],['b2','c1','c2'])
tree('beskarning-3',node,edges,v,['r','a'],['b2','c1','c2'])
# Parameter trace: alpha and beta are persistent at each active internal node.
b={'r':('−∞','+∞'),'a':('−∞','+∞'),'b':('−∞','+∞'),'c':('−∞','+∞')};v={};tree('alpha-0',node,edges,v,bounds=b)
frames=[({'a0':8},['a0'],[],{'a':('−∞',8)}),({'a1':5},['a1'],[],{'a':('−∞',5)}),({'a2':6,'a':5},['a','a2'],[],{}),({},['r','a'],[],{'r':(5,'+∞'),'b':(5,'+∞'),'c':(5,'+∞')}),({'b0':9},['b0'],[],{'b':(5,9)}),({'b1':3,'b':'≤3'},['b','b1'],['b2'],{'b':(5,3)}),({'c0':2,'c':'≤2'},['c','c0'],['b2','c1','c2'],{'c':(5,2)}),({'r':5},['r','a'],['b2','c1','c2'],{})]
for i,(add,act,pr,bs) in enumerate(frames,1):v.update(add);b.update(bs);tree(f'alpha-{i}',node,edges,v,act,pr,b,count=f'{sum(k in v for k in lk)} / 9 löv undersökta')
# Single MAX choice from source slide11.
n={'r':(900,180,'max'),'a':(460,580,'min'),'b':(900,580,'min'),'c':(1340,580,'min')};es=[('r',k) for k in 'abc'];tree('val-0',n,es,dict(zip('abc',[4,3,2])));tree('val-1',n,es,dict(zip('rabc',[4,4,3,2])),['r','a'])
# Original exercise18/19, including the approved correction of two values.
n={'r':(900,120,'max'),'a':(460,310,'min'),'b':(1340,310,'min')};es=[('r','a'),('r','b')]
for key,x,parent in [('aa',250,'a'),('ab',650,'a'),('ba',1150,'b'),('bb',1550,'b')]:
 n[key]=(x,500,'max');es.append((parent,key))
 for j,dx in enumerate([-75,75]):n[key+str(j)]=(x+dx,710,'min');es.append((key,key+str(j)))
v=dict(zip([p+str(j) for p in ['aa','ab','ba','bb'] for j in range(2)],[5,3,8,2,7,1,2,3]));tree('ovning-0',n,es,v)
for i,(add,act,pr) in enumerate([({'aa':5},['aa','aa0','aa1'],[]),({'ab':'≥8'},['ab','ab0'],['ab1']),({'a':5},['a','aa'],['ab1']),({'ba':7},['ba','ba0','ba1'],['ab1']),({'bb':3},['bb','bb0','bb1'],['ab1']),({'b':3},['b','bb'],['ab1']),({'r':5},['r','a'],['ab1'])],1):v.update(add);tree(f'ovning-{i}',n,es,v,act,pr)
# Slide20: derive exact adjacency from original connectors, then spread leaves for readability.
ds=DATA[20];groups=[d for d in ds if d['type']=='group'];ns={};es=[]
for g in groups:
 kids=[d for d in ds if d.get('parent')==g['id']];tri=next((d for d in kids if d.get('geom')=='triangle'),None)
 if not tri:continue
 val=next((d['text'] for d in kids if d.get('text')), '');ns[str(g['id'])]={'x':g['x']+g['w']/2,'y':g['y']+g['h']/2,'role':'min' if abs(tri['rot']-180)<1 else 'max','value':val}
for d in ds:
 if d['type']!='line':continue
 x1=d['x']+(d['w'] if d['fh'] else 0);x2=d['x']+(0 if d['fh'] else d['w']);y1=d['y']+(d['h'] if d['fv'] else 0);y2=d['y']+(0 if d['fv'] else d['h'])
 a=min(ns,key=lambda k:(ns[k]['x']-x1)**2+(ns[k]['y']-y1)**2);bb=min(ns,key=lambda k:(ns[k]['x']-x2)**2+(ns[k]['y']-y2)**2)
 if a!=bb:
  if ns[a]['y']>ns[bb]['y']:a,bb=bb,a
  es.append((a,bb))
# The source has one free text value at the left child of the right branch.
for d in ds:
 if d.get('text')=='5' and d.get('parent') is None:
  k=min(ns,key=lambda k:(ns[k]['x']-d['x']-d['w']/2)**2+(ns[k]['y']-d['y']-d['h']/2)**2);ns[k]['value']='5'
assert len(ns)==22 and len(set(es))==21,(len(ns),len(set(es)))
children={k:[] for k in ns}
for a,bb in es:children[a].append(bb)
for k in children:children[k].sort(key=lambda q:ns[q]['x'])
root=next(k for k in ns if k not in [b for a,b in es]);positions={};leafindex=[0]
def lay(k,depth):
 if not children[k]:x=leafindex[0];leafindex[0]+=1
 else:
  xs=[lay(c,depth+1) for c in children[k]];x=sum(xs)/len(xs)
 positions[k]=(x,depth,ns[k]['role']);return x
lay(root,0);maxdepth=max(p[1] for p in positions.values());nn={k:(100+x*1600/(leafindex[0]-1),95+dep*620/maxdepth,role) for k,(x,dep,role) in positions.items()}
tree('rakna-noder',nn,es,{k:v['value'] for k,v in ns.items()})
(ROOT/'docs/minimax/diagramdata-v2.json').write_text(json.dumps({'source20':{'nodes':ns,'edges':es,'root':root},'source12':{'leaves':[8,5,6,9,3,7,2,4,6],'root':5},'source19':{'leaves':[5,3,8,2,7,1,2,3],'root':5}},ensure_ascii=False,indent=2)+'\n')
print('Generated',len(list(OUT.glob('*.svg'))),'faithful diagram states; source20 has',len(ns),'nodes and',len(es),'edges.')
